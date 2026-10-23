"""Generador local IGP. Uso: python generar_dashboards.py fuente.xlsx --salida entrega --corte 2026-10-22
Dependencia de administración: openpyxl. Consulta HTML: sin dependencias.
"""
import argparse, collections, csv, datetime as dt, hashlib, json, re, unicodedata
from pathlib import Path
import openpyxl

REQUIRED = '''id_registro fila_sgcc tipo tipo_proceso macroproceso proceso subproceso_actividad partes_interesadas organo_unidad_organica descripcion aplica_sgc sgc_censis sgc_cenvul sgsi sgas sstt factor_riesgo cuestiones_contexto requisito_necesidades expectativas_partes_interesadas consecuencia_efecto clasificacion personal_expuesto probabilidad impacto puntaje_inicial nivel_inicial_reportado nivel_2025_reportado nivel_2026_reportado nivel_2026_validado inconsistencia_nivel_2026 sin_evaluacion_2026 tratamiento_reportado controles_existentes acciones_2024 acciones_2025 soporte plan_accion_vinculado responsable_acciones plan_contingencia documentado_registro tipo_control frecuencia_control modalidad_control existe_evidencia responsable_ejecucion_seguimiento puntaje_control resultado_control tipo_seguimiento fecha_implementacion fecha_seguimiento evaluacion_seguimiento fecha_cumplimiento eficacia_2026 eficacia_2024 referencia_origen'''.split()
BOOL = 'aplica_sgc sgc_censis sgc_cenvul sgsi sgas sstt inconsistencia_nivel_2026 sin_evaluacion_2026'.split()
def norm(s):
    return ''.join(c for c in unicodedata.normalize('NFD', str(s or '')) if unicodedata.category(c) != 'Mn').lower().strip()
def clean(v):
    return v.isoformat() if isinstance(v,(dt.date,dt.datetime)) else v
def table(w,name):
    rows=list(w[name].values); head=rows[0]
    if len(set(head)) != len(head) or None in head: raise ValueError('Encabezados inválidos: '+name)
    return [dict(zip(head,map(clean,r))) for r in rows[1:] if any(v is not None for v in r)]
def source(path):
    w=openpyxl.load_workbook(path,data_only=True)
    needed='REGISTROS ACCIONES_2025 VALIDACIONES CRITERIOS_MAPA DICCIONARIO CATALOGOS SGCC_ORIG CRITERIO_VF_ORIG GUIA'.split()
    if set(needed)-set(w.sheetnames): raise ValueError('Faltan hojas requeridas')
    r=table(w,'REGISTROS'); a=table(w,'ACCIONES_2025'); v=table(w,'VALIDACIONES')
    if not r or set(REQUIRED)-set(r[0]): raise ValueError('Faltan campos esenciales: '+str(set(REQUIRED)-set(r[0] if r else [])))
    ids=[x['id_registro'] for x in r]
    if any(not i for i in ids) or len(set(ids))!=len(ids): raise ValueError('IDs vacíos o duplicados')
    for x in r:
        if x['tipo'] not in ('RIESGO','OPORTUNIDAD') or not x['proceso']: raise ValueError('Tipo/proceso inválido')
        for key in BOOL:
            if not isinstance(x[key],bool): raise ValueError(f'{x["id_registro"]}: {key} debe ser booleano')
    for name,rows in [('ACCIONES_2025',a),('VALIDACIONES',v)]:
        if any(x.get('id_registro') not in ids for x in rows): raise ValueError('Relación huérfana: '+name)
    formulas=openpyxl.load_workbook(path,data_only=False)
    for name in ('REGISTROS','ACCIONES_2025','VALIDACIONES','CRITERIOS_MAPA'):
        for row in formulas[name]:
            for c in row:
                if c.data_type=='f' and w[name][c.coordinate].value is None: raise ValueError('Fórmula sin resultado: '+name+'!'+c.coordinate)
    return r,a,v,table(w,'CRITERIOS_MAPA'),table(w,'DICCIONARIO')
def reference_checks(r,a):
    sgc=[x for x in r if x['aplica_sgc']]; risks=[x for x in sgc if x['tipo']=='RIESGO']
    actual={'registros':len(r),'riesgos':sum(x['tipo']=='RIESGO' for x in r),'oportunidades':sum(x['tipo']=='OPORTUNIDAD' for x in r),'sgc':len(sgc),'riesgos_sgc':len(risks),'oportunidades_sgc':sum(x['tipo']=='OPORTUNIDAD' for x in sgc),'censis':sum(x['sgc_censis'] for x in r),'cenvul':sum(x['sgc_cenvul'] for x in r),'segmentos_auxiliares':len(a)}
    for key,label in [('bajos','riesgo bajo'),('moderados','riesgo moderado'),('altos','riesgo alto'),('criticos','riesgo critico')]: actual[key]=sum(norm(x['nivel_2026_validado'])==label for x in risks)
    expected=dict(registros=112,riesgos=95,oportunidades=17,sgc=99,riesgos_sgc=86,oportunidades_sgc=13,censis=77,cenvul=63,segmentos_auxiliares=202,bajos=30,moderados=55,altos=1,criticos=0)
    return [{'control':k,'real':actual[k],'referencia':n,'estado':'OK' if actual[k]==n else 'DISCREPANCIA — prevalece la fuente real'} for k,n in expected.items()]
def code(p):
    slug=re.sub('[^a-z0-9]+','-',norm(p)).strip('-')[:26]
    return slug+'-'+hashlib.sha256(p.encode()).hexdigest()[:6]
def safe_json(data): return json.dumps(data,ensure_ascii=False,separators=(',',':')).replace('<','\\u003c').replace('>','\\u003e').replace('&','\\u0026')
def strings(obj):
    if isinstance(obj,str): yield obj
    elif isinstance(obj,dict):
        for x in obj.values(): yield from strings(x)
    elif isinstance(obj,list):
        for x in obj: yield from strings(x)
def generate(path,out,cutoff,process=None,unit=None):
    dt.date.fromisoformat(cutoff)
    if out.exists() and any(out.iterdir()): raise ValueError('La carpeta de salida ya contiene archivos. Use una carpeta nueva para versionar la entrega.')
    r,a,v,criteria,dictionary=source(path)
    checks=reference_checks(r,a)
    out.mkdir(parents=True,exist_ok=True)
    template=Path(__file__).with_name('dashboard_template.html').read_text(encoding='utf-8')
    meta={'corte':cutoff,'generacion':dt.datetime.now().astimezone().isoformat(timespec='seconds'),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'fuente':path.name}
    labels={x['campo']:x['etiqueta_ui'] for x in dictionary}
    provenance={x['campo']:{'referencia':x['referencia_sgcc'],'interpretacion':x['interpretacion'],'advertencia':x['advertencia']} for x in dictionary}
    index=[]; isolation=[]
    def emit(rows,scope,rel,master=False):
        ids={x['id_registro'] for x in rows}
        model={'meta':{**meta,'ambito':scope,'maestro':master,'registros':len(rows)},'records':rows,'actions':[x for x in a if x['id_registro'] in ids],'validations':[x for x in v if x['id_registro'] in ids],'criteria':criteria,'labels':labels,'provenance':provenance}
        serialized=safe_json(model); content=template.replace('__PAYLOAD__',serialized)
        recovered=json.loads(re.search(r'<script id="data" type="application/json">(.*?)</script>',content,re.S)[1])
        assert {x['id_registro'] for x in recovered['records']}==ids
        assert all(x['id_registro'] in ids for x in recovered['actions']+recovered['validations'])
        assert recovered['records']==rows, 'Pérdida de texto original'
        if not master:
            forbidden=[x for x in r if x['id_registro'] not in ids]
            for x in forbidden:
                assert '"'+x['id_registro']+'"' not in serialized, 'ID ajeno'
                for key in ('descripcion','acciones_2025'):
                    val=x[key]
                    if isinstance(val,str) and len(val)>50 and not any(val in t for t in strings(model)):
                        assert safe_json(val)[1:-1] not in serialized, 'Texto ajeno'
        target=out/rel; target.parent.mkdir(parents=True,exist_ok=True); target.write_text(content,encoding='utf-8')
        index.append({'ambito':scope,'registros':len(rows),'riesgos':sum(x['tipo']=='RIESGO' for x in rows),'oportunidades':sum(x['tipo']=='OPORTUNIDAD' for x in rows),'registros_sgc':sum(x['aplica_sgc'] for x in rows),'archivo':rel})
        isolation.append({'archivo':rel,'ids_autorizados':sorted(ids),'integridad_textos':'OK','relaciones_y_aislamiento':'OK'})
    emit(r,'Institucional · OPPM / administración SIG','salidas/administracion/Dashboard_IGP_Institucional.html',True)
    for p in sorted({x['proceso'] for x in r}):
        c=code(p); emit([x for x in r if x['proceso']==p],p,f'salidas/auditados/{c}/Dashboard_Auditoria_{c}.html')
    if unit or process:
        subset=[x for x in r if (not unit or x['organo_unidad_organica']==unit) and (not process or x['proceso']==process)]
        if not subset: raise ValueError('Ámbito opcional sin coincidencias exactas')
        scope=' / '.join(x for x in (process,unit) if x); emit(subset,scope,'salidas/ambito_seleccionado/Dashboard_Auditoria_Ambito.html')
    with (out/'INDICE_DE_ENTREGA.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(index[0])); writer.writeheader(); writer.writerows(index)
    report={'metadatos':meta,'controles':checks,'archivos':isolation,'advertencias':['Los nombres volcanes/VOLCANES se conservan como procesos distintos.','Las fechas se contrastan con el corte explícito.','La suma de marcas no equivale al total único.']}
    (out/'VALIDACIONES_GENERACION.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'archivos':len(index),'controles':checks},ensure_ascii=False,indent=2))
    return report
if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__); parser.add_argument('fuente',type=Path); parser.add_argument('--salida',type=Path,required=True); parser.add_argument('--corte',required=True); parser.add_argument('--proceso'); parser.add_argument('--unidad'); args=parser.parse_args()
    try: generate(args.fuente,args.salida,args.corte,args.proceso,args.unidad)
    except Exception as e: raise SystemExit('ERROR DE GENERACIÓN: '+str(e))
